import openpyxl

def TPI(Codingcolumn,size):
    coding = 0
    tech_Productivity = 0
    for i in range(6,size+1):
        coding += (ws.cell(row=i,column=Codingcolumn)).value
    size = size -5
    tech_Productivity = coding/size
    return tech_Productivity
    
def AAI(Studycolumn,Classcolumn,size):
    studeTime = 0
    ClassTime = 0
    Academic_Activity = 0
    for i in range(6,size+1):
        studeTime += (ws.cell(row=i,column=Studycolumn)).value
        ClassTime += (ws.cell(row=i,column=Classcolumn)).value
    size = size -5
    Academic_Activity = (studeTime+ClassTime)/size
    return Academic_Activity

def PhAI(FitnessColumn,size):
    FitnessTime = 0
    Physical_Activity = 0
    for i in range(6,size+1):
        FitnessTime += (ws.cell(row=i,column=FitnessColumn)).value   
    size = size - 5
    Physical_Activity = FitnessTime/size
    return Physical_Activity
    
def SRI(SleepColumn, size):
    SleepTime = 0
    Sleep_Recovery = 0
    for i in range(6,size+1):
        SleepTime += (ws.cell(row=i,column=SleepColumn)).value   
    size = size - 5
    Sleep_Recovery = SleepTime/size
    return Sleep_Recovery

def ABI(UnaccountedColumn, size):
    UnaccountedTime = 0
    Activity_Balance = 0
    for i in range(6,size+1):
        UnaccountedTime += (ws.cell(row=i,column=UnaccountedColumn)).value  
    size = size - 5
    Activity_Balance = UnaccountedTime/size
    return Activity_Balance

def TUI(TotalTrackedColumn,size):
    TotalTrackedTime = 0
    Time_Utilization = 0
    for i in range(6,size+1):
        TotalTrackedTime += (ws.cell(row=i,column=TotalTrackedColumn)).value  
    size = size - 5
    Time_Utilization = TotalTrackedTime/size
    return Time_Utilization

def EI(DaysFellingColumn,SatisfactionLevelColumn,EnergyLevelColumn,size):
    DaysFellingIndex = 0
    SatisfactionLevelIndex = 0
    EnergyLevelIndex = 0
    Experience_Index = 0
    for i in range(6,size+1):
        DaysFellingValue = (ws.cell(row=i,column=DaysFellingColumn)).value
        if DaysFellingValue ==  "Excellent":
            DaysFellingIndex += 5
        elif DaysFellingValue ==  "Good":
            DaysFellingIndex += 4
        elif DaysFellingValue ==  "Neutral":
            DaysFellingIndex += 3
        elif DaysFellingValue ==  "Low":
            DaysFellingIndex += 2
        else:
            DaysFellingIndex += 1

        SatisfactionLevelValue = (ws.cell(row=i,column=SatisfactionLevelColumn)).value
        if SatisfactionLevelValue ==  "Very Satisfied":
            SatisfactionLevelIndex += 5
        elif SatisfactionLevelValue ==  "Satisfied":
            SatisfactionLevelIndex += 4
        elif SatisfactionLevelValue ==  "Neutral":
            SatisfactionLevelIndex += 3
        elif SatisfactionLevelValue ==  "Unsatisfied":
            SatisfactionLevelIndex += 2
        else:
            SatisfactionLevelIndex += 1  

        EnergyLevelValue = (ws.cell(row=i,column=EnergyLevelColumn)).value  
        if EnergyLevelValue ==  "High":
            EnergyLevelIndex += 3
        elif EnergyLevelValue ==  "Medium":
            EnergyLevelIndex += 2
        else:
            EnergyLevelIndex += 1
    size = size-5
    Experience_Index = (DaysFellingIndex+SatisfactionLevelIndex+EnergyLevelIndex)/(3*size)
    return Experience_Index

def DCI(size):
    size = size-5
    Data_Continuity = (size/40)*100
    return Data_Continuity

def PAI(TPI,AAI,PhAI,SRI,TUI,EI,DCI):
    PAI_Value = (0.15*TPI)+(AAI*0.20)+(0.15*PhAI)+(0.20*SRI)+(0.15*TUI)+(0.10*EI)+(0.05*DCI)
    return PAI_Value

wb = openpyxl.load_workbook("12612199.xlsx",data_only=True)
ws = wb["Daily Log"]
size = ws.max_row
TechProductivity = TPI(5,size)
AcademicActivity = AAI(4,6,size)
PhysicalActivity = PhAI(3,size)
SleepRecovery = SRI(2,size)
ActivityBalance = ABI(10,size)
TimeUtilization = TUI(9,size)
ExperienceIndex = EI(11,12,13,size)
DataCountinuity = DCI(size)
PersonalActivityIndex = PAI(TechProductivity,AcademicActivity,PhysicalActivity,SleepRecovery,TimeUtilization,ExperienceIndex,DataCountinuity)

print(f"My Data My Story:\nTech Productivity: {TechProductivity:.2f}\nAcademic Activity: {AcademicActivity:.2f}\nPhysical Activity: {PhysicalActivity:.2f}\nSleep Recovery: {SleepRecovery:.2f}\nActivity Balance: {ActivityBalance:.2f}\nTime Utilization: {TimeUtilization:.2f}\nExperience Index: {ExperienceIndex:.2f}\nData Countinuity : {DataCountinuity}\nPersonal Activity Index: {PersonalActivityIndex:.2f}")
wb.close()